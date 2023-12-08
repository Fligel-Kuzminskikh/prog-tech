from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from re import findall


class WorkerMsOffice:
    extensions = ["docx", "pptx", "xlsx", "xlsm", "xltm", "xltx"]
    filepath = str()

    def get_docx(self):
        metadata = dict()
        try:
            docx = Document(docx=self.filepath)
            properties = docx.core_properties
            for directory in dir(properties):
                if not directory.startswith('_'):
                    metadata[directory] = str(getattr(properties, directory))
        except:
            pass
        return metadata

    def get_xl(self):
        metadata = dict()
        try:
            workbook = load_workbook(self.filepath)
            for directory in dir(workbook.properties):
                if not directory.startswith('_') and directory not in ["to_tree", "from_tree"]:
                    metadata[directory] = str(getattr(workbook.properties, directory))
        except:
            pass
        return metadata

    def get_pptx(self):
        metadata = dict()
        try:
            pptx = Presentation(pptx=self.filepath)
            properties = pptx.core_properties
            for directory in dir(properties):
                if not directory.startswith('_'):
                    metadata[directory] = str(getattr(properties, directory))
        except:
            pass
        return metadata

    def get(self):
        if self.filepath:
            if findall(pattern=r"\.docx$", string=self.filepath):
                self.get_docx()
            elif findall(pattern=r"\.pptx$", string=self.filepath):
                self.get_pptx()
            elif findall(pattern=r"\.xl[st][xm]$", string=self.filepath):
                self.get_xl()
